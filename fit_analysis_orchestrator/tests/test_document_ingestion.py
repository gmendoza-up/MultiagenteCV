# -*- coding: utf-8 -*-
from __future__ import annotations

import asyncio
import hashlib
import os
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest

from fit_analysis_orchestrator.document_ingestion import DocumentIngestionService
from fit_analysis_orchestrator.models import DocumentIngestionResult


def write_text_file(path: Path, content: str) -> None:
    path.write_text(content, encoding="utf-8")


def test_ingest_text_file(tmp_path: Path):
    file_path = tmp_path / "candidate_a.txt"
    write_text_file(file_path, "Nombre: Juan Pérez\nExperiencia: Python")

    service = DocumentIngestionService()
    results = service.ingest_directory(str(tmp_path))

    assert len(results) == 1
    result = results[0]
    assert result.status == "processed"
    assert result.file_type == "txt"
    assert "Juan Pérez" in result.content
    assert result.page_count == 0
    assert result.error is None
    assert result.document_hash


def test_ingest_unsupported_file_type(tmp_path: Path):
    file_path = tmp_path / "candidate_a.xlsx"
    write_text_file(file_path, "datos")

    service = DocumentIngestionService()
    results = service.ingest_directory(str(tmp_path))

    assert len(results) == 1
    assert results[0].status == "failed"
    assert "no soportada" in results[0].error.lower()


def test_ingest_duplicate_files(tmp_path: Path):
    file_path1 = tmp_path / "candidate_a.txt"
    file_path2 = tmp_path / "candidate_b.txt"
    write_text_file(file_path1, "duplicado")
    write_text_file(file_path2, "duplicado")

    service = DocumentIngestionService()
    results = service.ingest_directory(str(tmp_path))

    assert len(results) == 2
    assert results[0].status == "processed"
    assert results[1].status == "failed"
    assert results[1].error == "duplicate document"


def test_ingest_partial_document_directory_skips_corrupt_file(tmp_path: Path):
    file_path1 = tmp_path / "candidate_a.txt"
    file_path2 = tmp_path / "candidate_b.pdf"
    write_text_file(file_path1, "contenido válido")
    file_path2.write_bytes(b"%PDF-1.4\n%%EOF\n\x00\x00")

    service = DocumentIngestionService()
    results = service.ingest_directory(str(tmp_path))

    assert len(results) == 2
    assert any(result.status == "processed" for result in results)
    assert any(result.status == "failed" for result in results)


@pytest.mark.skipif(True, reason="PDF and DOCX integration test requires optional dependencies")
def test_ingest_pdf_and_docx_files(tmp_path: Path):
    import fpdf
    from docx import Document

    pdf_path = tmp_path / "candidate_a.pdf"
    pdf = fpdf.FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(0, 10, "Nombre: Test PDF", ln=True)
    pdf.output(str(pdf_path))

    docx_path = tmp_path / "candidate_b.docx"
    document = Document()
    document.add_paragraph("Nombre: Test DOCX")
    document.save(str(docx_path))

    service = DocumentIngestionService()
    results = service.ingest_directory(str(tmp_path))

    assert len(results) == 2
    assert any(result.file_type == "pdf" and result.page_count == 1 for result in results)
    assert any(result.file_type == "docx" and "Test DOCX" in result.content for result in results)
